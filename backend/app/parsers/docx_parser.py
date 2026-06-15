import io
import docx
from app.parsers.base import BaseParser

class DocxParser(BaseParser):
    def parse(self, file_content: bytes) -> str:
        """
        Extracts plain text from DOCX binary content using python-docx.
        """
        try:
            file_stream = io.BytesIO(file_content)
            doc = docx.Document(file_stream)
            
            text_lines = []
            
            # Extract text from paragraphs
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    text_lines.append(paragraph.text)
            
            # Extract text from tables
            for table in doc.tables:
                for row in table.rows:
                    row_text = []
                    for cell in row.cells:
                        if cell.text.strip():
                            row_text.append(cell.text.strip())
                    if row_text:
                        text_lines.append(" | ".join(row_text))
                        
            return "\n".join(text_lines).strip()
        except Exception as e:
            raise ValueError(f"Failed to parse DOCX file: {str(e)}")
